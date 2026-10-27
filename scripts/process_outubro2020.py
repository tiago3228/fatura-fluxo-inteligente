from __future__ import annotations

from datetime import datetime
from pathlib import Path
import csv
import os
import sys
from openpyxl import load_workbook, Workbook

PRIVATE_DIR = Path.home() / ".local/share/fatura-fluxo-inteligente-private"
SOURCE = Path(
    os.environ.get(
        "FATURA_SOURCE_FILE",
        sys.argv[1] if len(sys.argv) > 1 else str(PRIVATE_DIR / "originals/OUTUBRODE2020.xlsx"),
    )
)
OUT_DIR = Path(os.environ.get("FATURA_IMPORT_OUT_DIR", str(PRIVATE_DIR / "imports")))
if not SOURCE.is_file():
    raise SystemExit(
        "Arquivo de origem não encontrado. Passe o caminho como argumento ou em FATURA_SOURCE_FILE."
    )
OUT_DIR.mkdir(parents=True, exist_ok=True)

PEOPLE = {
    "ATENA": "Atena Cardoso",
    "ARLETE": "Arlete",
    "ANA": "Ana Kamilly",
    "JOAO": "João Cardoso",
    "REU": "Reu",
}


def norm(value: object) -> str:
    text = str(value or "").strip().upper()
    return text.replace("Á", "A").replace("À", "A").replace("Ã", "A").replace("Â", "A").replace("É", "E").replace("Ê", "E").replace("Í", "I").replace("Ó", "O").replace("Ô", "O").replace("Õ", "O").replace("Ú", "U").replace("Ç", "C")


def date_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value or "").strip()


def row_record(sheet: str, row_number: int, person: str, date: object, description: object, amount: object, source_name: object, note: str = "") -> dict[str, object]:
    return {
        "planilha": sheet,
        "linha": row_number,
        "pessoa_atribuida": person,
        "data": date_text(date),
        "descricao": str(description or "").strip(),
        "valor": amount if amount is not None else "",
        "nome_original": str(source_name or "").strip(),
        "observacao": note,
    }


wb = load_workbook(SOURCE, data_only=True)
assigned: list[dict[str, object]] = []
unassigned: list[dict[str, object]] = []
excluded: list[dict[str, object]] = []

# Planilha2 contains cash movements and totals, not itemized purchases. Keep every
# person-labelled movement for review instead of inventing a purchase description.
ws = wb["Planilha2"]
for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
    name, date, amount = row[0], row[1], row[2]
    key = norm(name)
    if key == "ARLETE 1800/ATENA":
        unassigned.append(row_record(ws.title, row_number, "", date, "Movimentação com divisão indicada como ARLETE 1800/ATENA", amount, name, "Duas pessoas aparecem, mas a planilha não informa uma divisão completa e inequívoca do valor."))
    elif key == "JOAO" and row[8]:
        unassigned.append(row_record(ws.title, row_number, "", date, "Movimentação de JOÃO com anotação adicional", amount, name, f"A anotação adicional é: {row[8]}. Não foi criada uma divisão automática."))
    elif key == "CARAO":
        unassigned.append(row_record(ws.title, row_number, "", date, "Registro identificado como CARÃO", amount, name, "Não existe correspondência segura entre CARÃO e uma pessoa cadastrada."))
    elif key in PEOPLE and date is not None and amount is not None:
        unassigned.append(row_record(ws.title, row_number, "", date, "Movimentação com nome, mas sem item de compra", amount, name, f"O nome corresponde a {PEOPLE[key]}, mas esta aba não informa um item; não foi convertida em compra."))

# Planilha2 (4): these are person totals without date/item, so they are not imported as purchases.
ws = wb["Planilha2 (4)"]
for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
    name, date_or_blank, amount = row[0], row[1], row[2]
    key = norm(name)
    if key in PEOPLE and amount is not None:
        unassigned.append(row_record(ws.title, row_number, "", date_or_blank, "Resumo por pessoa sem data e sem item", amount, name, "Pode ser conferido como resumo, mas não foi convertido em compra individual."))
    elif key == "TOTAL GERAL":
        excluded.append(row_record(ws.title, row_number, "", date_or_blank, "TOTAL GERAL", amount, name, "Totais não são compras individuais."))

# ABRIL DE 2026: explicit person + item + amount rows are safe assignments.
ws = wb["ABRIL DE 2026"]
for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
    name, item, amount = row[0], row[1], row[2]
    key = norm(name)
    if key in PEOPLE and item and amount is not None:
        assigned.append(row_record(ws.title, row_number, PEOPLE[key], "", item, amount, name))
    elif name is None and item and amount is not None:
        excluded.append(row_record(ws.title, row_number, "", "", item, amount, name, "Resumo/fechamento sem coluna de pessoa; mantido fora das atribuições."))

# Export assigned and unresolved records.
headers = ["planilha", "linha", "pessoa_atribuida", "data", "descricao", "valor", "nome_original", "observacao"]
for filename, rows in [("OUTUBRODE2020_atribuidas.csv", assigned), ("OUTUBRODE2020_nao_atribuidas.csv", unassigned), ("OUTUBRODE2020_excluidas_controle.csv", excluded)]:
    with (OUT_DIR / filename).open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)

wb_out = Workbook()
wb_out.remove(wb_out.active)
for title, rows in [("Atribuídas", assigned), ("Não atribuídas", unassigned), ("Controle excluído", excluded)]:
    ws_out = wb_out.create_sheet(title)
    ws_out.append(headers)
    for record in rows:
        ws_out.append([record[h] for h in headers])
    ws_out.freeze_panes = "A2"
    ws_out.auto_filter.ref = ws_out.dimensions
    for column in ws_out.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        ws_out.column_dimensions[column[0].column_letter].width = width
wb_out.save(OUT_DIR / "OUTUBRODE2020_atribuicoes.xlsx")

print(f"Atribuídas: {len(assigned)}")
print(f"Não atribuídas/revisar: {len(unassigned)}")
print(f"Excluídas por serem totais/controle: {len(excluded)}")
print("Por pessoa:")
from collections import Counter
for name, count in Counter(str(row["pessoa_atribuida"]) for row in assigned).most_common():
    print(f"- {name}: {count}")
